import io
from pathlib import Path

import pytest
from docx import Document
from fastapi import UploadFile
from reportlab.pdfgen import canvas

from app.core.config import Settings
from app.services.candidate_extraction_service import CandidateExtractionService, ResumeExtractionError, assess_text_quality, detect_sections
from app.services.llm_service import LLMService,MockLLMProvider

BASE="""Jordan Lee
Software Engineer
SUMMARY
Builds reliable data products for fictional organizations.
EXPERIENCE
Software Engineer at Northstar Systems
Built a Python service and documented measured reliability improvements.
EDUCATION
B.S. Computer Science, Example University
"""

CASES=[
"simple_one_column","two_column","two_page_engineering","long_executive_cv","academic_cv","student_resume",
"research_scientist","product_manager","public_health","finance","tables","docx_tables","hyperlinks","unicode",
"icons_bullets","unusual_headings","no_skills","no_summary","multiple_roles","overlapping_dates","projects_heavy",
"sparse","very_short","long_resume","txt_resume","docx_resume","pdf_text_boxes","image_only","publications_patents",
"certifications_leadership",
]

def cfg(path:Path):return Settings(candidate_private_storage_path=str(path/"private"),candidate_profile_path=str(path/"profile.json"),candidate_evidence_path=str(path/"evidence.json"),candidate_preferences_path=str(path/"preferences.json"),llm_enabled=False)
def upload(name,data,mime):return UploadFile(filename=name,file=io.BytesIO(data),headers={"content-type":mime})
def service(path):return CandidateExtractionService(LLMService(MockLLMProvider(),cfg(path)),cfg(path))
def make_pdf(text):
    stream=io.BytesIO();pdf=canvas.Canvas(stream);y=760
    for line in text.splitlines():
        pdf.drawString(50,y,line[:110]);y-=14
        if y<50:pdf.showPage();y=760
    pdf.save();return stream.getvalue()
def make_docx(text,table=False):
    doc=Document();doc.add_paragraph(text)
    if table:
        grid=doc.add_table(rows=2,cols=2);grid.cell(0,0).text="Skills";grid.cell(0,1).text="Python, SQL";grid.cell(1,0).text="Projects";grid.cell(1,1).text="Search evaluation"
    stream=io.BytesIO();doc.save(stream);return stream.getvalue()

@pytest.mark.parametrize("case",CASES)
def test_fictional_resume_compatibility(case,tmp_path):
    text=BASE
    if case in {"long_executive_cv","academic_cv","long_resume"}:text += ("PUBLICATIONS\nFictional paper on dependable systems.\n"*90)
    if case=="student_resume":text=text.replace("EXPERIENCE","PROJECTS")
    if case=="research_scientist":text += "RESEARCH\nEvaluated fictional retrieval benchmarks.\n"
    if case=="product_manager":text=text.replace("Software Engineer","Product Manager")
    if case=="public_health":text=text.replace("Software Engineer","Public Health Analyst")
    if case=="finance":text=text.replace("Software Engineer","Financial Analyst")
    if case in {"unicode","icons_bullets"}:text += "SKILLS\n• Python · naïve Bayes · résumé analytics ✓\n"
    if case=="unusual_headings":text=text.replace("EXPERIENCE","CAREER HISTORY")
    if case=="no_summary":text=text.replace("SUMMARY\nBuilds reliable data products for fictional organizations.\n","")
    if case in {"multiple_roles","overlapping_dates"}:text += "Senior Engineer at Northstar Systems, 2022–Present\nEngineer at Northstar Systems, 2021–2023\n"
    if case=="projects_heavy":text += "SELECTED PROJECTS\n"+("Built a fictional test system.\n"*20)
    if case in {"sparse","very_short"}:text="Jordan Lee\nStudent developer with Python projects and coursework."
    if case=="publications_patents":text += "PATENTS\nFictional retrieval indexing method.\n"
    if case=="certifications_leadership":text += "CERTIFICATIONS\nCloud Fundamentals\nLEADERSHIP\nLed a student coding club.\n"
    if case=="image_only":
        data=make_pdf("");name="resume.pdf";mime="application/pdf"
        with pytest.raises(ResumeExtractionError):service(tmp_path).ingest(upload(name,data,mime))
        assert list((tmp_path/"private"/"resumes").glob("*.pdf"));return
    if case in {"docx_tables","docx_resume","tables","hyperlinks"}:data=make_docx(text,table=case in {"docx_tables","tables"});name="resume.docx";mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    elif case in {"txt_resume","very_short"}:data=text.encode("utf-8-sig");name="resume.txt";mime="text/plain"
    else:data=make_pdf(text);name="resume.pdf";mime="application/pdf"
    result=service(tmp_path).ingest(upload(name,data,mime))
    assert result.metadata.text_length>20 and result.metadata.text_quality_status in {"GOOD","USABLE","LOW_QUALITY"}
    assert (tmp_path/"private"/"resumes"/f"{result.metadata.document_id}.txt").exists()

def test_section_detection_and_quality_are_nonfatal_helpers():
    sections=detect_sections("ABOUT\nFictional builder\nWORK HISTORY\nEngineer at Example")
    assert {"summary","experience"}<=set(sections)
    assert assess_text_quality("",2)=="EMPTY"
    assert assess_text_quality("x",2)=="SCANNED_OR_IMAGE_ONLY"
